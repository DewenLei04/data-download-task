#!/usr/bin/env python3
"""Behavioral test for v4 multi-sheet coverage and novelty enforcement."""

from __future__ import annotations

import json
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook


SKILL = Path(__file__).resolve().parents[1]
SCRIPTS = SKILL / "scripts"
sys.path.insert(0, str(SCRIPTS))

from dataset_io import load_dataset  # noqa: E402
from validate_collections import novelty_metrics  # noqa: E402


def run(*args: str) -> None:
    subprocess.run([sys.executable, *args], check=True, capture_output=True, text=True)


def representation(path: str) -> dict:
    return {
        "generation_mode": "direct",
        "record_order": {
            "policy": "match_source",
            "generation_guidance": "Emit identifiers directly in final sequential source order.",
        },
        "field_profiles": [
            {
                "path": path,
                "value_role": "surrogate_identifier",
                "cardinality": "unique",
                "sequence_policy": "match_source",
                "record_layout": "preserve",
                "linked_paths": [],
                "direct_generation_guidance": "Generate a new unique sequential identifier directly in each final row.",
            }
        ],
    }


def generation_collection(path: str, identifier: str, value: str, exempt_status=False) -> dict:
    exemptions = []
    if exempt_status:
        exemptions.append(
            {"path": "/status", "reason": "The source defines a closed two-value vocabulary."}
        )
    fields = [
        {
            "path": identifier,
            "meaning": "synthetic row identifier",
            "physical_type": "string",
            "nullable": False,
            "unique": True,
            "generation_guidance": "Generate new values with the observed syntax.",
            "corruption_prohibited": True,
        },
        {
            "path": value,
            "meaning": "generated measurement",
            "physical_type": "integer",
            "nullable": False,
            "minimum": 0,
            "generation_guidance": "Generate inside the observed scale.",
        },
    ]
    if exempt_status:
        fields.append(
            {
                "path": "/status",
                "meaning": "closed status",
                "physical_type": "string",
                "nullable": False,
                "allowed_values": ["open", "closed"],
                "generation_guidance": "Reuse the controlled vocabulary.",
            }
        )
    return {
        "record_path": path,
        "record_count": 20,
        "primary_identifiers": [identifier],
        "grouping_keys": [],
        "protected_fields": [identifier],
        "representation": representation(identifier),
        "novelty": {
            "minimum_changed_value_ratio": 0.5,
            "high_similarity_threshold": 0.9,
            "maximum_high_similarity_record_ratio": 0.2,
            "exemptions": exemptions,
        },
        "hard_constraints": ["Identifiers are unique and measurements are nonnegative."],
        "soft_constraints": [],
        "fields": fields,
        "machine_rules": [],
    }


def main() -> None:
    with tempfile.TemporaryDirectory() as temp:
        root = Path(temp)
        source = root / "source.xlsx"
        clean = root / "clean.xlsx"
        dirty = root / "dirty.xlsx"
        report = root / "structure.json"
        spec_path = root / "spec.json"
        plan_path = root / "plan.json"
        patches_path = root / "patches.json"
        manifest_path = root / "manifest.json"
        validation_path = root / "validation.json"

        workbook = Workbook()
        alpha = workbook.active
        alpha.title = "Alpha"
        alpha.append(["id", "status", "value"])
        for index in range(1, 21):
            alpha.append([f"A{index:03d}", "open" if index % 2 else "closed", index * 10])
        beta = workbook.create_sheet("Beta")
        beta.append(["code", "amount"])
        for index in range(1, 21):
            beta.append([f"B{index:03d}", index * 100])
        docs = workbook.create_sheet("Read me")
        docs.append(["Workbook instructions"])
        docs.append(["This sheet is documentation and must remain unchanged."])
        workbook.save(source)

        generated = load_workbook(source)
        for index in range(2, 22):
            generated["Alpha"].cell(index, 1).value = f"A{index + 99:03d}"
            generated["Alpha"].cell(index, 3).value = index * 11
            generated["Beta"].cell(index, 1).value = f"B{index + 99:03d}"
            generated["Beta"].cell(index, 2).value = index * 111
        generated.save(clean)

        spec = {
            "version": 4,
            "source": {
                "physical_format": "excel",
                "collections": [
                    {"record_path": "sheet:Alpha", "record_count": 20, "role": "event rows", "handling": "generate"},
                    {"record_path": "sheet:Beta", "record_count": 20, "role": "amount rows", "handling": "generate"},
                    {"record_path": "sheet:Read me", "record_count": 1, "role": "workbook documentation", "handling": "preserve_documentation", "reason": "The sheet contains instructions, not sampled data."},
                ],
            },
            "scenario_inference": {
                "domain": "test operations",
                "dataset_role": "multi-table workbook fixture",
                "record_meaning": "one generated test record",
                "producer": "automated test workflow",
                "curation_level": "medium",
                "confidence": 1.0,
                "evidence": ["Two data sheets and one documentation sheet are directly observable."],
                "assumptions": [],
            },
            "generation": {
                "record_count": 40,
                "random_seed": 1,
                "source_model": {
                    "default_source_reuse": "allow",
                    "dimensions": [
                        {
                            "name": "two_table_generation",
                            "scope": "dataset",
                            "paths": [],
                            "evidence": "Both data sheets have identifiers and numeric values.",
                            "constraint_kind": "hard",
                            "generation_strategy": "bounded_extension",
                            "generation_rule": "Generate new identifiers and nearby nonnegative values in both data sheets.",
                        }
                    ],
                },
                "collections": [
                    generation_collection("sheet:Alpha", "/id", "/value", True),
                    generation_collection("sheet:Beta", "/code", "/amount"),
                ],
            },
            "dirty_policy": {
                "unit": "record",
                "max_rate": 0.05,
                "selected_rate": 0.05,
                "collection_coverage": {"mode": "all_eligible_when_feasible"},
                "rationale": "A small automated fixture uses the maximum permitted benchmark rate.",
                "error_types": [
                    {
                        "name": "negative_measurement",
                        "weight": 1.0,
                        "eligible_record_paths": ["sheet:Alpha", "sheet:Beta"],
                        "eligible_fields": ["/value", "/amount"],
                        "constraint_violated": "Measurements must be nonnegative.",
                    }
                ],
            },
            "quality": {
                "source_row_overlap_policy": "reject",
                "source_identifier_overlap_policy": "reject",
                "require_exact_record_count": True,
            },
        }
        spec_path.write_text(json.dumps(spec, indent=2) + "\n", encoding="utf-8")

        run(str(SCRIPTS / "inspect_dataset.py"), str(source), "--output", str(report))
        inspected = json.loads(report.read_text(encoding="utf-8"))
        assert [item["record_path"] for item in inspected["collections"]] == [
            "sheet:Alpha",
            "sheet:Beta",
            "sheet:Read me",
        ]

        run(str(SCRIPTS / "plan_corruption.py"), str(clean), str(spec_path), "--output", str(plan_path))
        plan = json.loads(plan_path.read_text(encoding="utf-8"))
        assert {item["record_path"] for item in plan["targets"]} == {"sheet:Alpha", "sheet:Beta"}
        assert plan["collection_coverage"]["all_eligible_coverage_feasible"] is True
        assert set(plan["collection_coverage"]["covered_record_paths"]) == {
            "sheet:Alpha",
            "sheet:Beta",
        }
        patches = {
            "version": 2,
            "patches": [
                {
                    "patch_id": item["patch_id"],
                    "record_path": item["record_path"],
                    "record_index": item["record_index"],
                    "error_type": item["error_type"],
                    "changes": [{"path": item["suggested_field"], "value": -1}],
                    "explanation": "A nonnegative measurement was replaced with a negative value.",
                }
                for item in plan["targets"]
            ],
        }
        patches_path.write_text(json.dumps(patches, indent=2) + "\n", encoding="utf-8")
        run(
            str(SCRIPTS / "apply_patches.py"),
            str(clean),
            str(patches_path),
            "--plan",
            str(plan_path),
            "--output",
            str(dirty),
            "--manifest",
            str(manifest_path),
        )
        run(
            str(SCRIPTS / "validate_collections.py"),
            "--source",
            str(source),
            "--clean",
            str(clean),
            "--dirty",
            str(dirty),
            "--spec",
            str(spec_path),
            "--manifest",
            str(manifest_path),
            "--output",
            str(validation_path),
        )
        validation = json.loads(validation_path.read_text(encoding="utf-8"))
        assert validation["passed"] is True

        source_beta = load_dataset(source, "sheet:Beta")
        unchanged_beta = load_dataset(source, "sheet:Beta")
        passed, metrics, violations = novelty_metrics(
            source_beta,
            unchanged_beta,
            spec["generation"]["collections"][1]["novelty"],
        )
        assert passed is False
        assert metrics["changed_value_ratio"] == 0.0
        assert violations


if __name__ == "__main__":
    main()
