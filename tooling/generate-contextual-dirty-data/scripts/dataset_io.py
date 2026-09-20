#!/usr/bin/env python3
"""Shared structured-dataset I/O helpers for the skill scripts."""

from __future__ import annotations

import csv
import copy
import datetime as dt
import decimal
import hashlib
import io
import json
import math
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable


class DatasetError(ValueError):
    """Raised when a dataset cannot be processed safely."""


@dataclass
class Dataset:
    path: Path
    format_name: str
    root: Any
    records: list[dict[str, Any]]
    record_path: str
    metadata: dict[str, Any] = field(default_factory=dict)


FORMAT_EXTENSIONS = {
    ".csv": "csv",
    ".tsv": "tsv",
    ".json": "json",
    ".jsonl": "jsonl",
    ".ndjson": "jsonl",
    ".yaml": "yaml",
    ".yml": "yaml",
    ".xlsx": "excel",
    ".xlsm": "excel",
    ".parquet": "parquet",
    ".pq": "parquet",
}


def json_safe(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        if math.isfinite(value):
            return value
        return str(value)
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, (decimal.Decimal, Path)):
        return str(value)
    if isinstance(value, bytes):
        return value.hex()
    if isinstance(value, dict):
        return {str(key): json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [json_safe(item) for item in value]
    return str(value)


def read_json(path: str | Path) -> Any:
    with Path(path).open("r", encoding="utf-8") as handle:
        return json.load(handle)


def write_json(path: str | Path, value: Any) -> None:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("w", encoding="utf-8", newline="\n") as handle:
        json.dump(json_safe(value), handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def detect_format(path: str | Path) -> str:
    source = Path(path)
    extension = source.suffix.lower()
    if extension in FORMAT_EXTENSIONS:
        return FORMAT_EXTENSIONS[extension]

    with source.open("rb") as handle:
        prefix = handle.read(8192).lstrip()
    if prefix.startswith((b"{", b"[")):
        try:
            json.loads(prefix.decode("utf-8"))
            return "json"
        except (UnicodeDecodeError, json.JSONDecodeError):
            pass
    raise DatasetError(
        f"Unsupported file extension {extension!r}. Supported formats: "
        "CSV, TSV, JSON, JSONL, YAML, XLSX, XLSM, and Parquet."
    )


def _escape_pointer_token(token: str) -> str:
    return token.replace("~", "~0").replace("/", "~1")


def _unescape_pointer_token(token: str) -> str:
    return token.replace("~1", "/").replace("~0", "~")


def pointer_tokens(pointer: str) -> list[str]:
    if pointer == "":
        return []
    if not pointer.startswith("/"):
        raise DatasetError(f"JSON Pointer must be empty or start with '/': {pointer!r}")
    return [_unescape_pointer_token(token) for token in pointer[1:].split("/")]


def get_pointer(value: Any, pointer: str) -> Any:
    current = value
    for token in pointer_tokens(pointer):
        if isinstance(current, dict):
            if token not in current:
                raise DatasetError(f"Path {pointer!r} does not exist")
            current = current[token]
        elif isinstance(current, list):
            try:
                current = current[int(token)]
            except (ValueError, IndexError) as exc:
                raise DatasetError(f"Invalid list segment {token!r} in {pointer!r}") from exc
        else:
            raise DatasetError(f"Path {pointer!r} traverses a scalar value")
    return current


def set_pointer(value: Any, pointer: str, replacement: Any) -> Any:
    tokens = pointer_tokens(pointer)
    if not tokens:
        return replacement
    parent = value
    for token in tokens[:-1]:
        if isinstance(parent, dict):
            if token not in parent:
                raise DatasetError(f"Path {pointer!r} does not exist")
            parent = parent[token]
        elif isinstance(parent, list):
            try:
                parent = parent[int(token)]
            except (ValueError, IndexError) as exc:
                raise DatasetError(f"Invalid list segment {token!r} in {pointer!r}") from exc
        else:
            raise DatasetError(f"Path {pointer!r} traverses a scalar value")

    final = tokens[-1]
    if isinstance(parent, dict):
        if final not in parent:
            raise DatasetError(f"Path {pointer!r} does not exist")
        parent[final] = replacement
    elif isinstance(parent, list):
        try:
            parent[int(final)] = replacement
        except (ValueError, IndexError) as exc:
            raise DatasetError(f"Invalid list segment {final!r} in {pointer!r}") from exc
    else:
        raise DatasetError(f"Path {pointer!r} traverses a scalar value")
    return value


def leaf_paths(record: dict[str, Any]) -> list[str]:
    paths: list[str] = []

    def visit(value: Any, base: str) -> None:
        if isinstance(value, dict):
            if not value:
                paths.append(base)
                return
            for key, item in value.items():
                visit(item, f"{base}/{_escape_pointer_token(str(key))}")
            return
        paths.append(base)

    visit(record, "")
    return sorted(path for path in paths if path)


def structural_paths(record: dict[str, Any]) -> list[str]:
    paths: set[str] = set()

    def visit(value: Any, base: str) -> None:
        if isinstance(value, dict):
            for key, item in value.items():
                child = f"{base}/{_escape_pointer_token(str(key))}"
                paths.add(child)
                visit(item, child)
        elif isinstance(value, list):
            for item in value:
                visit(item, f"{base}/*")

    visit(record, "")
    return sorted(paths)


def field_name_structure(value: Any) -> Any:
    """Return field-name paths while ignoring values, scalar types, and list length."""
    paths: set[str] = set()

    def visit(item: Any, base: str) -> None:
        if isinstance(item, dict):
            for key, child in item.items():
                path = f"{base}/{_escape_pointer_token(str(key))}"
                paths.add(path)
                visit(child, path)
        elif isinstance(item, list):
            for child in item:
                visit(child, f"{base}/*")

    visit(value, "")
    return tuple(sorted(paths))


def normalized_record_fingerprint(record: dict[str, Any]) -> str:
    encoded = json.dumps(
        json_safe(record), ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _decode_text(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()
    encodings = ("utf-8-sig",) if raw.startswith(b"\xef\xbb\xbf") else ("utf-8",)
    for encoding in encodings:
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    try:
        return raw.decode("cp1252"), "cp1252"
    except UnicodeDecodeError as exc:
        raise DatasetError("The text file is not valid UTF-8 or CP1252") from exc


def _find_record_collections(root: Any) -> list[tuple[str, list[dict[str, Any]]]]:
    candidates: list[tuple[str, list[dict[str, Any]]]] = []

    def visit(value: Any, pointer: str, depth: int) -> None:
        if depth > 10:
            return
        if isinstance(value, list):
            if value and all(isinstance(item, dict) for item in value):
                candidates.append((pointer, value))
                return
            for index, item in enumerate(value[:20]):
                visit(item, f"{pointer}/{index}", depth + 1)
        elif isinstance(value, dict):
            for key, item in value.items():
                escaped = _escape_pointer_token(str(key))
                visit(item, f"{pointer}/{escaped}", depth + 1)

    visit(root, "", 0)
    return candidates


def _select_records(root: Any, requested_path: str | None) -> tuple[str, list[dict[str, Any]]]:
    if requested_path is not None:
        selected = get_pointer(root, requested_path)
        if isinstance(selected, dict) and requested_path == "":
            return requested_path, [selected]
        if not isinstance(selected, list) or not all(isinstance(item, dict) for item in selected):
            raise DatasetError(
                f"Record path {requested_path!r} must select an object or an array of objects"
            )
        return requested_path, selected

    if isinstance(root, list) and all(isinstance(item, dict) for item in root):
        return "", root
    candidates = _find_record_collections(root)
    if candidates:
        candidates.sort(key=lambda item: (len(item[1]), -item[0].count("/")), reverse=True)
        return candidates[0]
    if isinstance(root, dict):
        return "", [root]
    raise DatasetError("No logical record collection could be inferred")


def _excel_header_row(worksheet: Any, scan_limit: int = 50) -> int:
    """Find a table header without mistaking workbook title/notes for field names.

    The first nonempty row remains the default.  A later row is selected only when
    it is a substantially denser, unique, text-dominant row near the top of the
    sheet.  This supports common workbooks that place titles and explanatory notes
    above the real table while keeping ordinary one-row-header sheets unchanged.
    """
    rows: list[tuple[int, list[Any]]] = []
    for row in worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, scan_limit)):
        values = [cell.value for cell in row]
        if any(value not in (None, "") for value in values):
            rows.append((row[0].row, values))
    if not rows:
        raise DatasetError(f"Sheet {worksheet.title!r} has no header row")

    first_row_number, first_values = rows[0]
    first_nonempty = sum(value not in (None, "") for value in first_values)
    candidates: list[tuple[int, int]] = []
    for row_number, values in rows:
        nonempty = [value for value in values if value not in (None, "")]
        if not nonempty:
            continue
        text_values = [value.strip() for value in nonempty if isinstance(value, str)]
        normalized = [value.casefold() for value in text_values if value]
        text_ratio = len(text_values) / len(nonempty)
        unique_text = len(normalized) == len(set(normalized))
        if text_ratio >= 0.9 and unique_text:
            candidates.append((len(nonempty), row_number))

    if candidates:
        densest_count, densest_row = max(candidates, key=lambda item: (item[0], -item[1]))
        if densest_row == first_row_number or densest_count >= max(2, first_nonempty * 2):
            return densest_row
    return first_row_number


def load_dataset(path: str | Path, record_path: str | None = None) -> Dataset:
    source = Path(path)
    if not source.is_file():
        raise DatasetError(f"Input file does not exist: {source}")
    format_name = detect_format(source)

    if format_name in {"csv", "tsv"}:
        if record_path not in (None, ""):
            raise DatasetError("CSV and TSV record paths must be empty")
        text, encoding = _decode_text(source)
        sample = text[:65536]
        default_delimiter = "\t" if format_name == "tsv" else ","
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
            delimiter = dialect.delimiter
            quotechar = dialect.quotechar or '"'
        except csv.Error:
            delimiter = default_delimiter
            quotechar = '"'
        reader = csv.DictReader(io.StringIO(text, newline=""), delimiter=delimiter, quotechar=quotechar)
        if not reader.fieldnames:
            raise DatasetError("The delimited file has no header row")
        if len(reader.fieldnames) != len(set(reader.fieldnames)):
            raise DatasetError("Duplicate header names are not supported")
        records = [dict(row) for row in reader]
        line_terminator = "\r\n" if "\r\n" in text else "\n"
        return Dataset(
            source,
            format_name,
            records,
            records,
            "",
            {
                "encoding": encoding,
                "columns": list(reader.fieldnames),
                "delimiter": delimiter,
                "quotechar": quotechar,
                "line_terminator": line_terminator,
            },
        )

    if format_name == "json":
        text, encoding = _decode_text(source)
        root = json.loads(text)
        selected_path, records = _select_records(root, record_path)
        indent = 2 if "\n" in text else None
        return Dataset(
            source, format_name, root, records, selected_path, {"encoding": encoding, "indent": indent}
        )

    if format_name == "jsonl":
        if record_path not in (None, ""):
            raise DatasetError("JSONL record paths must be empty")
        text, encoding = _decode_text(source)
        records: list[dict[str, Any]] = []
        for line_number, line in enumerate(text.splitlines(), start=1):
            if not line.strip():
                continue
            item = json.loads(line)
            if not isinstance(item, dict):
                raise DatasetError(f"JSONL line {line_number} is not an object")
            records.append(item)
        return Dataset(source, format_name, records, records, "", {"encoding": encoding})

    if format_name == "yaml":
        try:
            import yaml  # type: ignore
        except ImportError as exc:
            raise DatasetError("YAML support requires PyYAML") from exc
        text, encoding = _decode_text(source)
        root = yaml.safe_load(text)
        selected_path, records = _select_records(root, record_path)
        return Dataset(source, format_name, root, records, selected_path, {"encoding": encoding})

    if format_name == "excel":
        try:
            import openpyxl  # type: ignore
        except ImportError as exc:
            raise DatasetError("Excel support requires openpyxl") from exc
        workbook = openpyxl.load_workbook(source, data_only=False, keep_vba=source.suffix.lower() == ".xlsm")
        if record_path is None:
            worksheet = max(workbook.worksheets, key=lambda sheet: max(sheet.max_row - 1, 0))
            selected_path = f"sheet:{worksheet.title}"
        else:
            if not record_path.startswith("sheet:"):
                raise DatasetError("Excel record paths must use sheet:SHEET_NAME")
            sheet_name = record_path[6:]
            if sheet_name not in workbook.sheetnames:
                raise DatasetError(f"Workbook has no sheet named {sheet_name!r}")
            worksheet = workbook[sheet_name]
            selected_path = record_path
        header_row = _excel_header_row(worksheet)
        headers = [cell.value for cell in worksheet[header_row]]
        while headers and headers[-1] in (None, ""):
            headers.pop()
        if not headers:
            raise DatasetError(f"Sheet {worksheet.title!r} has no named data columns")
        columns = [str(value) if value is not None else f"__blank_column_{index + 1}" for index, value in enumerate(headers)]
        if len(columns) != len(set(columns)):
            raise DatasetError("Duplicate Excel header names are not supported")
        records = []
        row_numbers = []
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, max_col=len(columns)),
            start=header_row + 1,
        ):
            values = [cell.value for cell in row]
            if all(value is None for value in values):
                continue
            records.append(dict(zip(columns, values)))
            row_numbers.append(row_number)
        return Dataset(
            source,
            format_name,
            workbook,
            records,
            selected_path,
            {
                "columns": columns,
                "header_row": header_row,
                "row_numbers": row_numbers,
                "sheet_names": workbook.sheetnames,
                "keep_vba": source.suffix.lower() == ".xlsm",
            },
        )

    if format_name == "parquet":
        if record_path not in (None, ""):
            raise DatasetError("Parquet record paths must be empty")
        try:
            import pyarrow.parquet as parquet  # type: ignore
        except ImportError as exc:
            raise DatasetError("Parquet support requires pyarrow") from exc
        table = parquet.read_table(source)
        records = table.to_pylist()
        if not all(isinstance(item, dict) for item in records):
            raise DatasetError("Parquet rows could not be represented as objects")
        return Dataset(
            source,
            format_name,
            table,
            records,
            "",
            {"columns": table.column_names, "arrow_schema": table.schema},
        )

    raise DatasetError(f"Unsupported format: {format_name}")


def discover_record_paths(path: str | Path) -> list[str]:
    """Return every candidate logical collection in a structured container."""
    source = Path(path)
    format_name = detect_format(source)
    if format_name == "excel":
        try:
            import openpyxl  # type: ignore
        except ImportError as exc:
            raise DatasetError("Excel support requires openpyxl") from exc
        workbook = openpyxl.load_workbook(source, read_only=True, data_only=False)
        return [f"sheet:{name}" for name in workbook.sheetnames]
    if format_name == "json":
        text, _ = _decode_text(source)
        root = json.loads(text)
    elif format_name == "yaml":
        try:
            import yaml  # type: ignore
        except ImportError as exc:
            raise DatasetError("YAML support requires PyYAML") from exc
        text, _ = _decode_text(source)
        root = yaml.safe_load(text)
    else:
        return [""]

    if isinstance(root, list) and all(isinstance(item, dict) for item in root):
        return [""]
    candidates = _find_record_collections(root)
    if candidates:
        return [path for path, _ in candidates]
    if isinstance(root, dict):
        return [""]
    raise DatasetError("No logical record collection could be inferred")


def save_collection_datasets(datasets: list[Dataset], output_path: str | Path) -> None:
    """Serialize several edited collections into one source-shaped container."""
    if not datasets:
        raise DatasetError("At least one collection is required")
    formats = {dataset.format_name for dataset in datasets}
    source_paths = {dataset.path.resolve() for dataset in datasets}
    output = Path(output_path)
    if len(formats) != 1 or len(source_paths) != 1:
        raise DatasetError("Collections must come from the same source container and format")
    if output.resolve() in source_paths:
        raise DatasetError("Refusing to overwrite the input dataset")
    if len(datasets) == 1:
        save_dataset(datasets[0], output)
        return

    format_name = datasets[0].format_name
    output.parent.mkdir(parents=True, exist_ok=True)
    if format_name == "excel":
        workbook = datasets[0].root
        for dataset in datasets:
            if not dataset.record_path.startswith("sheet:"):
                raise DatasetError("Every Excel collection must use sheet:SHEET_NAME")
            worksheet = workbook[dataset.record_path[6:]]
            columns = dataset.metadata["columns"]
            for record, row_number in zip(dataset.records, dataset.metadata["row_numbers"]):
                for column_number, column in enumerate(columns, start=1):
                    worksheet.cell(row=row_number, column=column_number).value = record[column]
        workbook.save(output)
        return

    if format_name in {"json", "yaml"}:
        base = datasets[0]
        base.root = copy.deepcopy(base.root)
        for dataset in datasets:
            replacement: Any
            if dataset.record_path == "" and isinstance(dataset.root, dict):
                if len(dataset.records) != 1:
                    raise DatasetError("A root object collection must contain one record")
                replacement = copy.deepcopy(dataset.records[0])
            else:
                replacement = copy.deepcopy(dataset.records)
            base.root = set_pointer(base.root, dataset.record_path, replacement)
        save_dataset(base, output)
        return

    raise DatasetError(
        f"Multiple logical collections are not supported for format {format_name!r}"
    )


def save_dataset(dataset: Dataset, output_path: str | Path) -> None:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.resolve() == dataset.path.resolve():
        raise DatasetError("Refusing to overwrite the input dataset")

    if dataset.format_name in {"csv", "tsv"}:
        columns = dataset.metadata["columns"]
        with output.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=columns,
                delimiter=dataset.metadata["delimiter"],
                quotechar=dataset.metadata["quotechar"],
                lineterminator=dataset.metadata["line_terminator"],
                extrasaction="raise",
            )
            writer.writeheader()
            for record in dataset.records:
                if any(isinstance(record.get(column), (dict, list)) for column in columns):
                    raise DatasetError("CSV and TSV values must remain scalar")
                writer.writerow(record)
        return

    if dataset.format_name == "json":
        with output.open("w", encoding="utf-8", newline="\n") as handle:
            json.dump(json_safe(dataset.root), handle, ensure_ascii=False, indent=dataset.metadata["indent"])
            handle.write("\n")
        return

    if dataset.format_name == "jsonl":
        with output.open("w", encoding="utf-8", newline="\n") as handle:
            for record in dataset.records:
                json.dump(json_safe(record), handle, ensure_ascii=False, separators=(",", ":"))
                handle.write("\n")
        return

    if dataset.format_name == "yaml":
        try:
            import yaml  # type: ignore
        except ImportError as exc:
            raise DatasetError("YAML support requires PyYAML") from exc
        with output.open("w", encoding="utf-8", newline="\n") as handle:
            yaml.safe_dump(dataset.root, handle, allow_unicode=True, sort_keys=False)
        return

    if dataset.format_name == "excel":
        worksheet = dataset.root[dataset.record_path[6:]]
        columns = dataset.metadata["columns"]
        for record, row_number in zip(dataset.records, dataset.metadata["row_numbers"]):
            for column_number, column in enumerate(columns, start=1):
                worksheet.cell(row=row_number, column=column_number).value = record[column]
        dataset.root.save(output)
        return

    if dataset.format_name == "parquet":
        try:
            import pyarrow as arrow  # type: ignore
            import pyarrow.parquet as parquet  # type: ignore
        except ImportError as exc:
            raise DatasetError("Parquet support requires pyarrow") from exc
        table = arrow.Table.from_pylist(dataset.records, schema=dataset.metadata["arrow_schema"])
        parquet.write_table(table, output)
        return

    raise DatasetError(f"Unsupported format: {dataset.format_name}")


def field_exists(record: dict[str, Any], pointer: str) -> bool:
    try:
        get_pointer(record, pointer)
        return True
    except DatasetError:
        return False


def unique_structural_paths(records: Iterable[dict[str, Any]]) -> list[str]:
    paths: set[str] = set()
    for record in records:
        paths.update(structural_paths(record))
    return sorted(paths)
